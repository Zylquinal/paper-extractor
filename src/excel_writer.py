import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from src.exceptions import ExcelWriteError

MAIN_SHEET = "Paper Analysis"
Q_SHEET = "Questionnaires"
V_SHEET = "Variables"
R_SHEET = "Main References"
H_SHEET = "Hypotheses"
REV_SHEET = "Review Details"
RQ_SHEET = "Research Questions"

MAIN_HEADERS = [
    "Title", "Publication Year", "Authors", "DOI",
    "Methodology & Research Object", "Sample",
    "Data Analysis", "Main Result",
    "Limit / Research Gap", "Implications",
    "Theory", "Q", "V", "R", "H", "S", "T",
]
Q_HEADERS = ["Paper Title", "Statement/Question", "Reference", "Variable"]
V_HEADERS = ["Paper Title", "Variable Name", "Result", "References"]
R_HEADERS = ["Paper Title", "Reference Title", "Authors", "Explanation"]
H_HEADERS = ["Paper Title", "Hypothesis Statement", "Result"]
REV_HEADERS = ["Paper Title", "Review Type", "Source Databases", "Time Span",
               "Papers Included", "Synthesis Method", "Quality Assessment"]
RQ_HEADERS = ["Paper Title", "Research Question"]

ALL_SHEETS = [MAIN_SHEET, Q_SHEET, V_SHEET, R_SHEET, H_SHEET, REV_SHEET, RQ_SHEET]

HEADER_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, size=11)
BODY_FONT = Font(name="Arial", size=11)
LINK_FONT = Font(name="Arial", size=11, color="0563C1", underline="single")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")

COLUMN_WIDTHS = {
    MAIN_SHEET: {"A": (20, 50), "B": (8, 12), "C": (20, 50), "D": (15, 40),
                 "E": (25, 65), "F": (20, 60), "G": (20, 50),
                 "H": (25, 65), "I": (25, 65), "J": (20, 60),
                 "K": (20, 40), "L": (3, 6), "M": (3, 6), "N": (3, 6),
                 "O": (3, 6), "P": (3, 6), "Q": (3, 6)},
    Q_SHEET: {"A": (20, 50), "B": (30, 80), "C": (15, 40), "D": (20, 50)},
    V_SHEET: {"A": (20, 50), "B": (20, 50), "C": (30, 80), "D": (15, 40)},
    R_SHEET: {"A": (20, 50), "B": (20, 50), "C": (20, 40), "D": (30, 80)},
    H_SHEET: {"A": (20, 50), "B": (40, 100), "C": (15, 30)},
    REV_SHEET: {"A": (20, 50), "B": (15, 40), "C": (20, 50),
                "D": (15, 30), "E": (10, 15), "F": (20, 50), "G": (20, 50)},
    RQ_SHEET: {"A": (20, 50), "B": (50, 120)},
}


def _style_header(ws, headers):
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_data_row(ws, row_num, num_cols):
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        if not cell.hyperlink:
            cell.font = BODY_FONT
        cell.border = THIN_BORDER
        cell.alignment = WRAP_ALIGNMENT


def _auto_width(ws, headers, widths):
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        min_w, max_w = widths.get(letter, (15, 60))
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        width = min(max(max_len + 2, min_w), max_w)
        ws.column_dimensions[letter].width = width


def _validate_headers(ws, headers):
    existing = []
    for col_idx in range(1, len(headers) + 1):
        existing.append(str(ws.cell(row=1, column=col_idx).value or ""))
    return existing == headers


def _validate_all_sheets(wb):
    sheets_headers = {
        MAIN_SHEET: MAIN_HEADERS,
        Q_SHEET: Q_HEADERS,
        V_SHEET: V_HEADERS,
        R_SHEET: R_HEADERS,
        H_SHEET: H_HEADERS,
        REV_SHEET: REV_HEADERS,
        RQ_SHEET: RQ_HEADERS,
    }
    for name in sheets_headers:
        if name not in wb.sheetnames:
            raise ExcelWriteError(
                f"Sheet '{name}' not found in workbook. "
                f"Expected sheets: {list(sheets_headers.keys())}"
            )
    for name, headers in sheets_headers.items():
        ws = wb[name]
        if not _validate_headers(ws, headers):
            raise ExcelWriteError(
                f"Sheet '{name}' has different header format. "
                f"Expected: {headers}"
            )


def _apply_link(ws, row, col, location, display):
    cell = ws.cell(row=row, column=col)
    cell.value = display
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=location)
    cell.font = LINK_FONT
    cell.alignment = Alignment(horizontal="center", vertical="top")


def _init_sheet(ws, headers):
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=header)
    _style_header(ws, headers)


def _format_authors(authors):
    if not authors:
        return ""
    if len(authors) == 1:
        return authors[0]
    return ", ".join(authors[:-1]) + " & " + authors[-1]


def _format_sample(samp):
    if not samp:
        return ""
    parts = []
    if samp.get("population"):
        parts.append(f"Populasi: {samp['population']}")
    if samp.get("size") is not None:
        parts.append(f"Ukuran: {samp['size']}")
    if samp.get("characteristics"):
        parts.append(f"Karakteristik: {samp['characteristics']}")
    return " | ".join(parts)


def _format_implications(imp):
    if not imp:
        return ""
    parts = []
    if imp.get("theoretical"):
        parts.append(f"Teoritis: {imp['theoretical']}")
    if imp.get("practical"):
        parts.append(f"Praktis: {imp['practical']}")
    return "\n".join(parts)


def _populate_main_sheet(ws, data, row, q_start, v_start, r_start, h_start,
                         s_start, t_start):
    theory = ""
    if data.get("model") and data.get("model", {}).get("theory"):
        theory = "; ".join(data["model"]["theory"])

    values = [
        data.get("title", ""),
        data.get("publication_year", ""),
        _format_authors(data.get("authors", [])),
        data.get("doi", ""),
        data.get("methodology_object", ""),
        _format_sample(data.get("sample")),
        "; ".join(data.get("data_analysis", [])),
        data.get("result", ""),
        data.get("limit_gap", ""),
        _format_implications(data.get("implications")),
        theory,
    ]
    for col, value in enumerate(values, 1):
        ws.cell(row=row, column=col, value=value)

    _style_data_row(ws, row, len(MAIN_HEADERS))

    if q_start:
        _apply_link(ws, row, 12, f"'{Q_SHEET}'!A{q_start}", "Q")
    if v_start:
        _apply_link(ws, row, 13, f"'{V_SHEET}'!A{v_start}", "V")
    if r_start:
        _apply_link(ws, row, 14, f"'{R_SHEET}'!A{r_start}", "R")
    if h_start:
        _apply_link(ws, row, 15, f"'{H_SHEET}'!A{h_start}", "H")
    if s_start:
        _apply_link(ws, row, 16, f"'{REV_SHEET}'!A{s_start}", "S")
    if t_start:
        _apply_link(ws, row, 17, f"'{RQ_SHEET}'!A{t_start}", "T")


def _write_backlink(ws, row, title, main_row):
    cell = ws.cell(row=row, column=1)
    cell.value = title
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{MAIN_SHEET}'!A{main_row}")
    cell.font = LINK_FONT
    cell.alignment = WRAP_ALIGNMENT


def _populate_questionnaires(ws, data, start_row, main_row):
    items = data.get("questionnaires")
    if not items:
        return
    title = data.get("title", "")
    for i, item in enumerate(items):
        row = start_row + i
        ref = item.get("reference")
        if ref is None:
            ref_text = ""
        elif ref == "UNKNOWN":
            ref_text = "UNKNOWN"
        else:
            ref_text = str(ref)
        var = item.get("variable")
        ws.cell(row=row, column=2, value=item.get("statement_question", ""))
        ws.cell(row=row, column=3, value=ref_text)
        ws.cell(row=row, column=4, value=var if var else "")
        _style_data_row(ws, row, len(Q_HEADERS))
        _write_backlink(ws, row, title, main_row)


def _populate_variables(ws, data, start_row, main_row):
    model = data.get("model")
    if not model or not model.get("variables"):
        return
    title = data.get("title", "")
    for i, var in enumerate(model["variables"]):
        row = start_row + i
        refs = var.get("references", [])
        ws.cell(row=row, column=2, value=var.get("name", ""))
        ws.cell(row=row, column=3, value=var.get("result", ""))
        if refs:
            ref_text = "; ".join(
                f"{r['author']} ({r['paper']})" if isinstance(r, dict) else str(r)
                for r in refs
            )
        else:
            ref_text = ""
        ws.cell(row=row, column=4, value=ref_text)
        _style_data_row(ws, row, len(V_HEADERS))
        _write_backlink(ws, row, title, main_row)


def _populate_references(ws, data, start_row, main_row):
    items = data.get("main_references")
    if not items:
        return
    title = data.get("title", "")
    for i, ref in enumerate(items):
        row = start_row + i
        ws.cell(row=row, column=2, value=ref.get("title", ""))
        ws.cell(row=row, column=3, value="; ".join(ref.get("authors", [])))
        ws.cell(row=row, column=4, value=ref.get("explanation", ""))
        _style_data_row(ws, row, len(R_HEADERS))
        _write_backlink(ws, row, title, main_row)


def _populate_hypotheses(ws, data, start_row, main_row):
    items = data.get("hypothesis")
    if not items:
        return
    title = data.get("title", "")
    for i, item in enumerate(items):
        row = start_row + i
        ws.cell(row=row, column=2, value=item.get("statement", ""))
        ws.cell(row=row, column=3, value=item.get("result", ""))
        _style_data_row(ws, row, len(H_HEADERS))
        _write_backlink(ws, row, title, main_row)


def _populate_review_details(ws, data, start_row, main_row):
    items = data.get("review_type")
    if not items:
        return
    title = data.get("title", "")
    row = start_row
    ws.cell(row=row, column=2, value=data.get("review_type", ""))
    ws.cell(row=row, column=3, value="; ".join(data.get("source_databases", [])))
    ts = data.get("time_span")
    if ts:
        ws.cell(row=row, column=4, value=f"{ts.get('start', '')}–{ts.get('end', '')}")
    ws.cell(row=row, column=5, value=data.get("papers_included", ""))
    ws.cell(row=row, column=6, value=data.get("synthesis_method", ""))
    ws.cell(row=row, column=7, value="; ".join(data.get("quality_assessment", [])))
    _style_data_row(ws, row, len(REV_HEADERS))
    _write_backlink(ws, row, title, main_row)


def _populate_research_questions(ws, data, start_row, main_row):
    items = data.get("research_questions")
    if not items:
        return
    title = data.get("title", "")
    for i, item in enumerate(items):
        row = start_row + i
        ws.cell(row=row, column=2, value=item)
        _style_data_row(ws, row, len(RQ_HEADERS))
        _write_backlink(ws, row, title, main_row)


def _auto_width_all(wb):
    for name, headers in [(MAIN_SHEET, MAIN_HEADERS), (Q_SHEET, Q_HEADERS),
                          (V_SHEET, V_HEADERS), (R_SHEET, R_HEADERS),
                          (H_SHEET, H_HEADERS), (REV_SHEET, REV_HEADERS),
                          (RQ_SHEET, RQ_HEADERS)]:
        _auto_width(wb[name], headers, COLUMN_WIDTHS[name])


def _create_workbook(data: dict) -> Workbook:
    wb = Workbook()

    ws_main = wb.active
    ws_main.title = MAIN_SHEET
    _init_sheet(ws_main, MAIN_HEADERS)

    ws_q = wb.create_sheet(Q_SHEET)
    _init_sheet(ws_q, Q_HEADERS)

    ws_v = wb.create_sheet(V_SHEET)
    _init_sheet(ws_v, V_HEADERS)

    ws_r = wb.create_sheet(R_SHEET)
    _init_sheet(ws_r, R_HEADERS)

    ws_h = wb.create_sheet(H_SHEET)
    _init_sheet(ws_h, H_HEADERS)

    ws_rev = wb.create_sheet(REV_SHEET)
    _init_sheet(ws_rev, REV_HEADERS)

    ws_rq = wb.create_sheet(RQ_SHEET)
    _init_sheet(ws_rq, RQ_HEADERS)

    main_row = 2

    has_q = bool(data.get("questionnaires"))
    has_v = bool(data.get("model") and data.get("model", {}).get("variables"))
    has_r = bool(data.get("main_references"))
    has_h = bool(data.get("hypothesis"))
    has_s = bool(data.get("review_type"))
    has_t = bool(data.get("research_questions"))

    q_start = 2 if has_q else None
    v_start = 2 if has_v else None
    r_start = 2 if has_r else None
    h_start = 2 if has_h else None
    s_start = 2 if has_s else None
    t_start = 2 if has_t else None

    _populate_main_sheet(ws_main, data, main_row, q_start, v_start, r_start, h_start,
                         s_start, t_start)

    if has_q:
        _populate_questionnaires(ws_q, data, q_start, main_row)
    if has_v:
        _populate_variables(ws_v, data, v_start, main_row)
    if has_r:
        _populate_references(ws_r, data, r_start, main_row)
    if has_h:
        _populate_hypotheses(ws_h, data, h_start, main_row)
    if has_s:
        _populate_review_details(ws_rev, data, s_start, main_row)
    if has_t:
        _populate_research_questions(ws_rq, data, t_start, main_row)

    _auto_width_all(wb)
    ws_main.freeze_panes = "A2"

    return wb


def write_result(data: dict, excel_path: str) -> None:
    if os.path.exists(excel_path):
        wb = load_workbook(excel_path)
        _validate_all_sheets(wb)

        ws_main = wb[MAIN_SHEET]
        main_row = ws_main.max_row + 1

        ws_q = wb[Q_SHEET]
        ws_v = wb[V_SHEET]
        ws_r = wb[R_SHEET]
        ws_h = wb[H_SHEET]
        ws_rev = wb[REV_SHEET]
        ws_rq = wb[RQ_SHEET]

        has_q = bool(data.get("questionnaires"))
        has_v = bool(data.get("model") and data.get("model", {}).get("variables"))
        has_r = bool(data.get("main_references"))
        has_h = bool(data.get("hypothesis"))
        has_s = bool(data.get("review_type"))
        has_t = bool(data.get("research_questions"))

        q_start = ws_q.max_row + 1 if has_q else None
        v_start = ws_v.max_row + 1 if has_v else None
        r_start = ws_r.max_row + 1 if has_r else None
        h_start = ws_h.max_row + 1 if has_h else None
        s_start = ws_rev.max_row + 1 if has_s else None
        t_start = ws_rq.max_row + 1 if has_t else None

        _populate_main_sheet(ws_main, data, main_row, q_start, v_start, r_start, h_start,
                             s_start, t_start)

        if has_q:
            _populate_questionnaires(ws_q, data, q_start, main_row)
        if has_v:
            _populate_variables(ws_v, data, v_start, main_row)
        if has_r:
            _populate_references(ws_r, data, r_start, main_row)
        if has_h:
            _populate_hypotheses(ws_h, data, h_start, main_row)
        if has_s:
            _populate_review_details(ws_rev, data, s_start, main_row)
        if has_t:
            _populate_research_questions(ws_rq, data, t_start, main_row)

        _auto_width_all(wb)
    else:
        wb = _create_workbook(data)

    wb.save(excel_path)

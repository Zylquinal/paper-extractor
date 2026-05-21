import os

import pytest
from openpyxl import load_workbook

from src.excel_writer import (
    MAIN_HEADERS, MAIN_SHEET, Q_HEADERS, Q_SHEET,
    V_HEADERS, V_SHEET, R_HEADERS, R_SHEET,
    H_HEADERS, H_SHEET, REV_HEADERS, REV_SHEET,
    RQ_HEADERS, RQ_SHEET, write_result,
)
from src.exceptions import ExcelWriteError


def test_create_new_excel(sample_extraction, temp_excel_path):
    if os.path.exists(temp_excel_path):
        os.unlink(temp_excel_path)

    write_result(sample_extraction, temp_excel_path)

    assert os.path.exists(temp_excel_path)
    wb = load_workbook(temp_excel_path)

    # All 7 sheets exist
    assert set(wb.sheetnames) == {MAIN_SHEET, Q_SHEET, V_SHEET, R_SHEET,
                                  H_SHEET, REV_SHEET, RQ_SHEET}

    # --- Main sheet ---
    ws_main = wb[MAIN_SHEET]
    assert ws_main.cell(row=1, column=1).value == "Title"
    assert ws_main.cell(row=1, column=2).value == "Publication Year"
    assert ws_main.cell(row=1, column=3).value == "Authors"
    assert ws_main.cell(row=1, column=4).value == "DOI"
    assert ws_main.cell(row=1, column=5).value == "Methodology & Research Object"
    assert ws_main.cell(row=1, column=6).value == "Sample"
    assert ws_main.cell(row=1, column=7).value == "Data Analysis"
    assert ws_main.cell(row=1, column=8).value == "Main Result"
    assert ws_main.cell(row=1, column=9).value == "Limit / Research Gap"
    assert ws_main.cell(row=1, column=10).value == "Implications"
    assert ws_main.cell(row=1, column=11).value == "Theory"
    assert ws_main.cell(row=1, column=12).value == "Q"
    assert ws_main.cell(row=1, column=13).value == "V"
    assert ws_main.cell(row=1, column=14).value == "R"
    assert ws_main.cell(row=1, column=15).value == "H"
    assert ws_main.cell(row=1, column=16).value == "S"
    assert ws_main.cell(row=1, column=17).value == "T"

    assert ws_main.cell(row=2, column=1).value == sample_extraction["title"]
    assert ws_main.cell(row=2, column=2).value == sample_extraction["publication_year"]
    assert ws_main.cell(row=2, column=3).value == "John Doe, Jane Smith & Bob Johnson"
    assert ws_main.cell(row=2, column=4).value == sample_extraction["doi"]
    assert ws_main.cell(row=2, column=5).value == sample_extraction["methodology_object"]
    assert ws_main.cell(row=2, column=6).value == "Populasi: Data kriminalitas di Kota Jakarta | Ukuran: 1500 | Karakteristik: Data historis 5 tahun terakhir, mencakup 10 kecamatan"
    assert ws_main.cell(row=2, column=7).value == "Random Forest; Gradient Boosting; Logistic Regression"
    assert ws_main.cell(row=2, column=8).value == sample_extraction["result"]
    assert ws_main.cell(row=2, column=9).value == sample_extraction["limit_gap"]
    assert ws_main.cell(row=2, column=10).value == "Teoritis: Memperkuat penerapan ensemble learning dalam prediksi kriminalitas temporal\nPraktis: Hasil dapat digunakan kepolisian untuk mengoptimalkan jadwal patroli"
    assert ws_main.cell(row=2, column=11).value == "Technology Acceptance Model"

    # Forward links
    q_cell = ws_main.cell(row=2, column=12)
    assert q_cell.value == "Q"
    assert q_cell.hyperlink is not None

    v_cell = ws_main.cell(row=2, column=13)
    assert v_cell.value == "V"
    assert v_cell.hyperlink is not None

    r_cell = ws_main.cell(row=2, column=14)
    assert r_cell.value == "R"
    assert r_cell.hyperlink is not None

    h_cell = ws_main.cell(row=2, column=15)
    assert h_cell.value == "H"
    assert h_cell.hyperlink is not None

    s_cell = ws_main.cell(row=2, column=16)
    assert s_cell.value is None  # No SLR data in empirical sample
    assert s_cell.hyperlink is None

    t_cell = ws_main.cell(row=2, column=17)
    assert t_cell.value is None
    assert t_cell.hyperlink is None

    # Header formatting
    header_cell = ws_main.cell(row=1, column=1)
    assert header_cell.font.bold is True
    assert header_cell.fill.start_color.rgb in ("DCE6F1", "00DCE6F1")

    assert ws_main.freeze_panes == "A2"

    # --- Q sheet ---
    ws_q = wb[Q_SHEET]
    assert ws_q.cell(row=1, column=1).value == Q_HEADERS[0]
    assert ws_q.cell(row=1, column=2).value == Q_HEADERS[1]
    assert ws_q.cell(row=1, column=3).value == Q_HEADERS[2]
    assert ws_q.cell(row=1, column=4).value == Q_HEADERS[3]

    assert ws_q.max_row == 3  # header + 2 items
    q_item0 = sample_extraction["questionnaires"][0]
    assert ws_q.cell(row=2, column=1).value == sample_extraction["title"]
    assert ws_q.cell(row=2, column=1).hyperlink is not None
    assert ws_q.cell(row=2, column=2).value == q_item0["statement_question"]
    assert ws_q.cell(row=2, column=3).value == q_item0["reference"]
    assert ws_q.cell(row=2, column=4).value == q_item0["variable"]

    q_item1 = sample_extraction["questionnaires"][1]
    assert ws_q.cell(row=3, column=1).value == sample_extraction["title"]
    assert ws_q.cell(row=3, column=2).value == q_item1["statement_question"]
    assert ws_q.cell(row=3, column=3).value is None  # reference is None → empty cell
    assert ws_q.cell(row=3, column=4).value == q_item1["variable"]

    # --- V sheet ---
    ws_v = wb[V_SHEET]
    assert ws_v.cell(row=1, column=1).value == V_HEADERS[0]
    assert ws_v.cell(row=1, column=2).value == V_HEADERS[1]
    assert ws_v.cell(row=1, column=3).value == V_HEADERS[2]
    assert ws_v.cell(row=1, column=4).value == V_HEADERS[3]

    assert ws_v.max_row == 3  # header + 2 variables
    v0 = sample_extraction["model"]["variables"][0]
    assert ws_v.cell(row=2, column=1).value == sample_extraction["title"]
    assert ws_v.cell(row=2, column=1).hyperlink is not None
    assert ws_v.cell(row=2, column=2).value == v0["name"]
    assert ws_v.cell(row=2, column=3).value == v0["result"]
    assert ws_v.cell(row=2, column=4).value == "Fred D. Davis (Perceived Usefulness, Perceived Ease of Use, and User Acceptance of Information Technology)"

    v1 = sample_extraction["model"]["variables"][1]
    assert ws_v.cell(row=3, column=2).value == v1["name"]
    assert ws_v.cell(row=3, column=4).value is None  # empty references → empty cell

    # --- R sheet ---
    ws_r = wb[R_SHEET]
    assert ws_r.cell(row=1, column=1).value == R_HEADERS[0]
    assert ws_r.cell(row=1, column=2).value == R_HEADERS[1]
    assert ws_r.cell(row=1, column=3).value == R_HEADERS[2]
    assert ws_r.cell(row=1, column=4).value == R_HEADERS[3]

    assert ws_r.max_row == 2  # header + 1 reference
    r0 = sample_extraction["main_references"][0]
    assert ws_r.cell(row=2, column=1).value == sample_extraction["title"]
    assert ws_r.cell(row=2, column=1).hyperlink is not None
    assert ws_r.cell(row=2, column=2).value == r0["title"]
    assert ws_r.cell(row=2, column=3).value == r0["authors"][0]
    assert ws_r.cell(row=2, column=4).value == r0["explanation"]

    # --- H sheet ---
    ws_h = wb[H_SHEET]
    assert ws_h.cell(row=1, column=1).value == H_HEADERS[0]
    assert ws_h.cell(row=1, column=2).value == H_HEADERS[1]
    assert ws_h.cell(row=1, column=3).value == H_HEADERS[2]

    assert ws_h.max_row == 3  # header + 2 hypotheses
    h0 = sample_extraction["hypothesis"][0]
    assert ws_h.cell(row=2, column=1).value == sample_extraction["title"]
    assert ws_h.cell(row=2, column=1).hyperlink is not None
    assert ws_h.cell(row=2, column=2).value == h0["statement"]
    assert ws_h.cell(row=2, column=3).value == h0["result"]

    h1 = sample_extraction["hypothesis"][1]
    assert ws_h.cell(row=3, column=1).value == sample_extraction["title"]
    assert ws_h.cell(row=3, column=1).hyperlink is not None
    assert ws_h.cell(row=3, column=2).value == h1["statement"]
    assert ws_h.cell(row=3, column=3).value == h1["result"]


def test_append_to_existing_excel(sample_extraction, temp_excel_path):
    if os.path.exists(temp_excel_path):
        os.unlink(temp_excel_path)

    write_result(sample_extraction, temp_excel_path)

    second = dict(sample_extraction)
    second["title"] = "Judul Kedua"
    write_result(second, temp_excel_path)

    wb = load_workbook(temp_excel_path)

    # Main sheet: 3 rows (header + 2 papers)
    ws_main = wb[MAIN_SHEET]
    assert ws_main.max_row == 3
    assert ws_main.cell(row=2, column=1).value == sample_extraction["title"]
    assert ws_main.cell(row=3, column=1).value == "Judul Kedua"
    assert ws_main.cell(row=1, column=1).value == MAIN_HEADERS[0]

    # Detail sheets: doubled row counts
    ws_q = wb[Q_SHEET]
    assert ws_q.max_row == 5  # header + 2+2 items
    # Paper 2 backlinks point to main row 3
    assert ws_q.cell(row=4, column=1).value == "Judul Kedua"
    assert ws_q.cell(row=4, column=1).hyperlink is not None

    ws_v = wb[V_SHEET]
    assert ws_v.max_row == 5  # header + 2+2 variables
    assert ws_v.cell(row=4, column=1).value == "Judul Kedua"
    assert ws_v.cell(row=4, column=1).hyperlink is not None

    ws_r = wb[R_SHEET]
    assert ws_r.max_row == 3  # header + 1+1 references
    assert ws_r.cell(row=3, column=1).value == "Judul Kedua"
    assert ws_r.cell(row=3, column=1).hyperlink is not None

    ws_h = wb[H_SHEET]
    assert ws_h.max_row == 5  # header + 2+2 hypotheses
    assert ws_h.cell(row=4, column=1).value == "Judul Kedua"
    assert ws_h.cell(row=4, column=1).hyperlink is not None

    # SLR/Reviews sheets have only headers (empirical data has no SLR fields)
    assert wb[REV_SHEET].max_row == 1
    assert wb[RQ_SHEET].max_row == 1


def test_append_multiple_results(sample_extraction, temp_excel_path):
    if os.path.exists(temp_excel_path):
        os.unlink(temp_excel_path)

    for i in range(3):
        data = dict(sample_extraction)
        data["title"] = f"Paper {i + 1}"
        write_result(data, temp_excel_path)

    wb = load_workbook(temp_excel_path)

    ws_main = wb[MAIN_SHEET]
    assert ws_main.max_row == 4  # header + 3 papers
    for i in range(3):
        assert ws_main.cell(row=i + 2, column=1).value == f"Paper {i + 1}"

    # With 3 papers each having 2 Q, 2 V, 1 R, 2 H
    assert wb[Q_SHEET].max_row == 7   # header + 3*2
    assert wb[V_SHEET].max_row == 7   # header + 3*2
    assert wb[R_SHEET].max_row == 4   # header + 3*1
    assert wb[H_SHEET].max_row == 7   # header + 3*2
    # SLR sheets empty (empirical data)
    assert wb[REV_SHEET].max_row == 1
    assert wb[RQ_SHEET].max_row == 1


def test_reject_mismatched_headers(tmp_path):
    mismatched = tmp_path / "wrong_headers.xlsx"

    from openpyxl import Workbook

    wb = Workbook()
    # Create all 7 expected sheets but with wrong headers
    ws = wb.active
    ws.title = MAIN_SHEET
    for i in range(17):
        ws.cell(row=1, column=i + 1, value=f"Wrong {i + 1}")

    for name in [Q_SHEET, V_SHEET, R_SHEET, H_SHEET, REV_SHEET, RQ_SHEET]:
        ws2 = wb.create_sheet(name)
        ws2.cell(row=1, column=1, value="Bad")

    wb.save(str(mismatched))

    with pytest.raises(ExcelWriteError, match="different header format"):
        write_result(
            {"title": "T", "methodology_object": "M", "result": "R", "limit_gap": "L"},
            str(mismatched),
        )


def test_reject_missing_sheets(tmp_path):
    only_main = tmp_path / "missing_sheets.xlsx"

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = MAIN_SHEET
    for i, h in enumerate(MAIN_HEADERS, 1):
        ws.cell(row=1, column=i, value=h)
    wb.save(str(only_main))

    with pytest.raises(ExcelWriteError, match="not found in workbook"):
        write_result(
            {"title": "T", "methodology_object": "M", "result": "R", "limit_gap": "L"},
            str(only_main),
        )


def test_no_optional_fields(sample_extraction_no_optional, temp_excel_path):
    if os.path.exists(temp_excel_path):
        os.unlink(temp_excel_path)

    write_result(sample_extraction_no_optional, temp_excel_path)

    wb = load_workbook(temp_excel_path)

    ws_main = wb[MAIN_SHEET]
    assert ws_main.cell(row=2, column=1).value == sample_extraction_no_optional["title"]
    assert ws_main.cell(row=2, column=2).value is None   # No publication year
    assert ws_main.cell(row=2, column=3).value is None   # No authors
    assert ws_main.cell(row=2, column=4).value is None   # No DOI
    assert ws_main.cell(row=2, column=6).value is None   # No sample
    assert ws_main.cell(row=2, column=7).value is None   # No data analysis
    assert ws_main.cell(row=2, column=10).value is None  # No implications
    assert ws_main.cell(row=2, column=11).value is None  # No theory
    assert ws_main.cell(row=2, column=12).value is None  # No Q link
    assert ws_main.cell(row=2, column=13).value is None  # No V link
    assert ws_main.cell(row=2, column=14).value is None  # No R link
    assert ws_main.cell(row=2, column=15).value is None  # No H link
    assert ws_main.cell(row=2, column=16).value is None  # No S link
    assert ws_main.cell(row=2, column=17).value is None  # No T link

    # Detail sheets: only headers, no data rows
    assert wb[Q_SHEET].max_row == 1
    assert wb[V_SHEET].max_row == 1
    assert wb[R_SHEET].max_row == 1
    assert wb[H_SHEET].max_row == 1
    assert wb[REV_SHEET].max_row == 1
    assert wb[RQ_SHEET].max_row == 1


def test_detail_hyperlink_target(sample_extraction, temp_excel_path):
    if os.path.exists(temp_excel_path):
        os.unlink(temp_excel_path)

    write_result(sample_extraction, temp_excel_path)

    wb = load_workbook(temp_excel_path)

    # Main sheet forward links point to correct detail rows
    ws_main = wb[MAIN_SHEET]
    q_hl = ws_main.cell(row=2, column=12).hyperlink
    assert q_hl is not None
    assert q_hl.location == f"'{Q_SHEET}'!A2"

    v_hl = ws_main.cell(row=2, column=13).hyperlink
    assert v_hl is not None
    assert v_hl.location == f"'{V_SHEET}'!A2"

    r_hl = ws_main.cell(row=2, column=14).hyperlink
    assert r_hl is not None
    assert r_hl.location == f"'{R_SHEET}'!A2"

    h_hl = ws_main.cell(row=2, column=15).hyperlink
    assert h_hl is not None
    assert h_hl.location == f"'{H_SHEET}'!A2"

    s_hl = ws_main.cell(row=2, column=16).hyperlink
    assert s_hl is None  # empirical sample has no SLR data

    t_hl = ws_main.cell(row=2, column=17).hyperlink
    assert t_hl is None

    # Detail sheet backlinks point to main sheet row 2
    q_back = wb[Q_SHEET].cell(row=2, column=1).hyperlink
    assert q_back is not None
    assert q_back.location == f"'{MAIN_SHEET}'!A2"

    v_back = wb[V_SHEET].cell(row=2, column=1).hyperlink
    assert v_back is not None
    assert v_back.location == f"'{MAIN_SHEET}'!A2"

    r_back = wb[R_SHEET].cell(row=2, column=1).hyperlink
    assert r_back is not None
    assert r_back.location == f"'{MAIN_SHEET}'!A2"

    h_back = wb[H_SHEET].cell(row=2, column=1).hyperlink
    assert h_back is not None
    assert h_back.location == f"'{MAIN_SHEET}'!A2"


def test_append_hyperlinks_target_correct_rows(sample_extraction, temp_excel_path):
    if os.path.exists(temp_excel_path):
        os.unlink(temp_excel_path)

    write_result(sample_extraction, temp_excel_path)

    second = dict(sample_extraction)
    second["title"] = "Judul Kedua"
    write_result(second, temp_excel_path)

    wb = load_workbook(temp_excel_path)

    # Paper 2 forward links point to detail rows for paper 2
    ws_main = wb[MAIN_SHEET]
    q_hl = ws_main.cell(row=3, column=12).hyperlink
    assert q_hl.location == f"'{Q_SHEET}'!A4"  # 1 header + 2 paper1 rows + 1 = row 4

    v_hl = ws_main.cell(row=3, column=13).hyperlink
    assert v_hl.location == f"'{V_SHEET}'!A4"

    r_hl = ws_main.cell(row=3, column=14).hyperlink
    assert r_hl.location == f"'{R_SHEET}'!A3"  # 1 header + 1 paper1 row + 1 = row 3

    h_hl = ws_main.cell(row=3, column=15).hyperlink
    assert h_hl.location == f"'{H_SHEET}'!A4"  # 1 header + 2 paper1 rows + 1 = row 4

    s_hl = ws_main.cell(row=3, column=16).hyperlink
    assert s_hl is None  # no SLR data

    t_hl = ws_main.cell(row=3, column=17).hyperlink
    assert t_hl is None

    # Paper 2 detail backlinks point to main row 3
    q_back = wb[Q_SHEET].cell(row=4, column=1).hyperlink
    assert q_back.location == f"'{MAIN_SHEET}'!A3"

    v_back = wb[V_SHEET].cell(row=4, column=1).hyperlink
    assert v_back.location == f"'{MAIN_SHEET}'!A3"

    r_back = wb[R_SHEET].cell(row=3, column=1).hyperlink
    assert r_back.location == f"'{MAIN_SHEET}'!A3"

    h_back = wb[H_SHEET].cell(row=4, column=1).hyperlink
    assert h_back.location == f"'{MAIN_SHEET}'!A3"


def test_slr_extraction(sample_extraction_slr, temp_excel_path):
    if os.path.exists(temp_excel_path):
        os.unlink(temp_excel_path)

    write_result(sample_extraction_slr, temp_excel_path)
    wb = load_workbook(temp_excel_path)

    assert set(wb.sheetnames) == {MAIN_SHEET, Q_SHEET, V_SHEET, R_SHEET,
                                  H_SHEET, REV_SHEET, RQ_SHEET}

    ws_main = wb[MAIN_SHEET]
    assert ws_main.cell(row=2, column=1).value == sample_extraction_slr["title"]
    assert ws_main.cell(row=2, column=12).value is None  # No Q link
    assert ws_main.cell(row=2, column=13).value is None  # No V link
    assert ws_main.cell(row=2, column=15).value is None  # No H link
    assert ws_main.cell(row=2, column=16).value == "S"   # S link
    assert ws_main.cell(row=2, column=16).hyperlink is not None
    assert ws_main.cell(row=2, column=17).value == "T"   # T link
    assert ws_main.cell(row=2, column=17).hyperlink is not None

    ws_rev = wb[REV_SHEET]
    assert ws_rev.cell(row=1, column=1).value == REV_HEADERS[0]
    assert ws_rev.cell(row=1, column=2).value == REV_HEADERS[1]
    assert ws_rev.cell(row=2, column=1).value == sample_extraction_slr["title"]
    assert ws_rev.cell(row=2, column=2).value == sample_extraction_slr["review_type"]
    assert ws_rev.cell(row=2, column=3).value == "Scopus; Web of Science; IEEE Xplore"
    assert ws_rev.cell(row=2, column=4).value == "2019–2024"
    assert ws_rev.cell(row=2, column=5).value == 87
    assert ws_rev.cell(row=2, column=6).value == "Thematic Synthesis"
    assert ws_rev.cell(row=2, column=7).value == "MMAT; CASP"

    ws_rq = wb[RQ_SHEET]
    assert ws_rq.cell(row=1, column=1).value == RQ_HEADERS[0]
    assert ws_rq.cell(row=1, column=2).value == RQ_HEADERS[1]
    assert ws_rq.max_row == 3  # header + 2 questions
    assert ws_rq.cell(row=2, column=1).value == sample_extraction_slr["title"]
    assert ws_rq.cell(row=2, column=2).value == sample_extraction_slr["research_questions"][0]
    assert ws_rq.cell(row=3, column=1).value == sample_extraction_slr["title"]
    assert ws_rq.cell(row=3, column=2).value == sample_extraction_slr["research_questions"][1]
